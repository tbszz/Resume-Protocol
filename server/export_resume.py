"""Render structured, verified resume data. stdin JSON -> stdout PDF/DOCX bytes."""
import io
import json
import os
import sys
from xml.sax.saxutils import escape

data = json.load(sys.stdin)
resume = data['resume']
sections = [('教育经历', 'education'), ('专业技能', 'skills'), ('项目经历', 'projects'), ('工作与实习经历', 'experience'), ('荣誉与证书', 'awards')]

def text(value):
    return str(value or '').replace('\x00', '')

contact = '  ·  '.join(text(v) for v in resume.get('contact', {}).values() if v)
heading = text(resume.get('targetTitle') or resume.get('title'))
output = io.BytesIO()
if data['format'] == 'pdf':
    from reportlab.pdfgen import canvas
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    font_path = os.environ.get('RESUME_PROTOCOL_FONT', 'C:/Windows/Fonts/msyh.ttc')
    if os.path.isfile(font_path):
        pdfmetrics.registerFont(TTFont('Resume', font_path, subfontIndex=0))
        font = 'Resume'
    else:
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        font = 'STSong-Light'
    body = ParagraphStyle('Body', fontName=font, fontSize=10, leading=17, textColor=colors.HexColor('#252525'), wordWrap='CJK', spaceAfter=7)
    title = ParagraphStyle('Title', parent=body, fontSize=23, leading=32, spaceAfter=6)
    sub = ParagraphStyle('Sub', parent=body, fontSize=10, textColor=colors.HexColor('#555555'))
    section = ParagraphStyle('Section', parent=body, fontSize=12, leading=20, spaceBefore=11, spaceAfter=7, borderWidth=0, borderPadding=3)
    def para(value, style=body):
        return Paragraph(escape(text(value)).replace('\n', '<br/>'), style)
    story = [para(resume.get('name') or '个人简历', title), para(heading, sub)]
    if resume.get('sourceNotice'): story.append(para(resume['sourceNotice'], sub))
    if contact: story.append(para(contact, sub))
    if resume.get('summary'): story.extend([Spacer(1, 8), para(resume['summary'])])
    for label, key in sections:
        items = [text(v) for v in resume.get(key, []) if v]
        if not items: continue
        if key == 'skills': items = [' · '.join(items)]
        story.append(KeepTogether([para(label, section), para(items[0])]))
        story.extend(para(v) for v in items[1:])
    def footer(c, doc):
        c.saveState()
        c.setFont(font, 8)
        c.setFillColor(colors.HexColor('#777777'))
        c.drawRightString(550, 25, str(doc.page))
        c.restoreState()
    SimpleDocTemplate(output, pagesize=(595.28,841.89), rightMargin=45, leftMargin=45, topMargin=38, bottomMargin=42, title=text(resume.get('name')) + ' - 简历', author='').build(story, onFirstPage=footer, onLaterPages=footer)
else:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.6)
    sec.left_margin = sec.right_margin = Inches(.65)
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.25
    doc.add_heading(text(resume.get('name') or '个人简历'), 0)
    doc.add_paragraph(heading)
    if resume.get('sourceNotice'): doc.add_paragraph(text(resume['sourceNotice']))
    if contact: doc.add_paragraph(contact)
    if resume.get('summary'): doc.add_paragraph(text(resume['summary']))
    for label, key in sections:
        items = [text(v) for v in resume.get(key, []) if v]
        if not items: continue
        doc.add_heading(label, 1)
        for value in ([' · '.join(items)] if key == 'skills' else items): doc.add_paragraph(value)
    doc.save(output)
sys.stdout.buffer.write(output.getvalue())
